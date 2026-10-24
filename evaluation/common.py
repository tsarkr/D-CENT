"""Paths, statistics, and annotation workbooks shared by the evaluation."""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any, Sequence

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

from pipeline.config import ROOT

EVAL_DIR = ROOT / "evaluation"
DATASET_DIR = EVAL_DIR / "dataset"
QUESTIONS = DATASET_DIR / "questions.json"        # every candidate question (written by `dataset --generate`)
BENCHMARK = DATASET_DIR / "benchmark.json"        # the questions that passed validation (used by `run`)
THIRD_PARTY_DIR = DATASET_DIR / "third_party"     # questions written outside the author team, merged on --generate
VALIDATION_DIR = DATASET_DIR / "validation"       # validation workbook and summary
RESULTS_DIR = EVAL_DIR / "results"                # one folder per run, plus graph/ (graph size, expert validation)
GRAPH_RESULTS = RESULTS_DIR / "graph"
ANNOTATION_DIR = EVAL_DIR / "annotation"          # workbooks for historians
EXPERT_DIR = ANNOTATION_DIR / "expert"            # graph validation workbooks
RATING_DIR = ANNOTATION_DIR / "rating"            # blind answer rating workbooks (hand this folder to raters)
RATING_KEY_DIR = ANNOTATION_DIR / "rating_key"    # blinding key (never give it to raters)


def latest_run() -> Path:
    """Most recent run folder under results/."""
    runs = sorted(p for p in RESULTS_DIR.glob("*") if (p / "scores.csv").exists())
    if not runs:
        raise SystemExit("평가 결과가 없습니다. 먼저 `python -m evaluation run`을 실행하세요.")
    return runs[-1]


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    text = json.dumps(data, ensure_ascii=False, indent=2)
    path.write_text(re.sub(r"\bNaN\b", "null", text), encoding="utf-8")  # NaN is not valid JSON


# ── Statistics (paired / resampling based, numpy only) ──────────────────────


def bootstrap_ci(values: Sequence[float], n_boot: int = 10000, alpha: float = 0.05, seed: int = 0) -> tuple[float, float, float]:
    """Return (mean, lower, upper) with a percentile bootstrap interval."""
    arr = np.asarray(values, dtype=float)
    if arr.size == 0:
        return float("nan"), float("nan"), float("nan")
    if arr.size == 1:
        return float(arr[0]), float(arr[0]), float(arr[0])
    rng = np.random.default_rng(seed)
    idx = rng.integers(0, arr.size, size=(n_boot, arr.size))
    means = arr[idx].mean(axis=1)
    lo, hi = np.quantile(means, [alpha / 2, 1 - alpha / 2])
    return float(arr.mean()), float(lo), float(hi)


def paired_permutation_test(a: Sequence[float], b: Sequence[float], n_perm: int = 20000, seed: int = 0) -> float:
    """Two-sided sign-flip permutation p-value for mean(a - b) on paired samples."""
    diff = np.asarray(a, dtype=float) - np.asarray(b, dtype=float)
    if diff.size == 0:
        return float("nan")
    if not np.any(diff):
        return 1.0
    observed = abs(diff.mean())
    rng = np.random.default_rng(seed)
    signs = rng.choice([-1.0, 1.0], size=(n_perm, diff.size))
    permuted = np.abs((signs * diff).mean(axis=1))
    # +1 correction keeps the Monte Carlo p-value strictly positive.
    return float((np.sum(permuted >= observed - 1e-12) + 1) / (n_perm + 1))


def holm_adjust(p_values: Sequence[float]) -> list[float]:
    """Holm-Bonferroni step-down adjustment (NaN entries are left untouched)."""
    adjusted = [float("nan")] * len(p_values)
    valid = sorted((p, i) for i, p in enumerate(p_values) if p == p)
    running = 0.0
    for rank, (p, i) in enumerate(valid):
        running = max(running, min(1.0, (len(valid) - rank) * p))
        adjusted[i] = running
    return adjusted


def prf(tp: int, fp: int, fn: int) -> tuple[float, float, float]:
    precision = tp / (tp + fp) if (tp + fp) else float("nan")
    recall = tp / (tp + fn) if (tp + fn) else float("nan")
    if precision != precision or recall != recall or (precision + recall) == 0:
        return precision, recall, float("nan")
    return precision, recall, 2 * precision * recall / (precision + recall)


def gwet_ac1(labels_a: Sequence, labels_b: Sequence) -> float:
    """Gwet's AC1 for two raters and nominal categories.

    Unlike Cohen's kappa it stays stable when one category dominates (most items judged correct),
    so it is reported next to kappa and raw agreement.
    """
    if len(labels_a) != len(labels_b):
        raise ValueError("두 평가자의 항목 수가 다릅니다.")
    n = len(labels_a)
    categories = sorted(set(labels_a) | set(labels_b))
    if n == 0 or len(categories) < 2:
        return float("nan")
    observed = sum(a == b for a, b in zip(labels_a, labels_b)) / n
    shares = [(list(labels_a).count(c) + list(labels_b).count(c)) / (2 * n) for c in categories]
    chance = sum(p * (1 - p) for p in shares) / (len(categories) - 1)
    return float("nan") if chance == 1 else (observed - chance) / (1 - chance)


def cohen_kappa(labels_a: Sequence, labels_b: Sequence, weights: str | None = None) -> float:
    """Cohen's kappa for two raters; ``weights='quadratic'`` for ordinal scales."""
    if len(labels_a) != len(labels_b):
        raise ValueError("두 평가자의 항목 수가 다릅니다.")
    if not labels_a:
        return float("nan")
    categories = sorted(set(labels_a) | set(labels_b))
    if len(categories) < 2:
        return float("nan")  # no variation: agreement is undefined, not perfect
    index = {c: i for i, c in enumerate(categories)}
    k = len(categories)
    observed = np.zeros((k, k))
    for x, y in zip(labels_a, labels_b):
        observed[index[x], index[y]] += 1
    observed /= observed.sum()
    expected = np.outer(observed.sum(axis=1), observed.sum(axis=0))
    if weights == "quadratic":
        positions = np.asarray([float(c) for c in categories])
        span = positions.max() - positions.min()
        disagreement = ((positions[:, None] - positions[None, :]) / span) ** 2
    else:
        disagreement = 1.0 - np.eye(k)
    denom = float((disagreement * expected).sum())
    if denom == 0:
        return float("nan")
    return 1.0 - float((disagreement * observed).sum()) / denom


# ── Annotation workbooks ─────────────────────────────────────────────────────
EXCEL_CELL_LIMIT = 32000


def add_sheet(
    wb: Workbook,
    title: str,
    columns: list[str],
    rows: list[dict[str, Any]],
    widths: dict[str, int] | None = None,
    choices: dict[str, list[str]] | None = None,
) -> None:
    """Append a sheet; ``choices`` maps a column to its allowed dropdown values."""
    ws = wb.create_sheet(title)
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([str(row.get(c, ""))[:EXCEL_CELL_LIMIT] if row.get(c) is not None else "" for c in columns])
    ws.freeze_panes = "A2"
    for idx, column in enumerate(columns, 1):
        letter = ws.cell(row=1, column=idx).column_letter
        ws.column_dimensions[letter].width = (widths or {}).get(column, 16)
        if (widths or {}).get(column, 0) >= 40:
            for cell in ws[letter][1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        if choices and column in choices:
            validation = DataValidation(type="list", formula1='"' + ",".join(choices[column]) + '"', allow_blank=True)
            ws.add_data_validation(validation)
            validation.add(f"{letter}2:{letter}{max(len(rows) + 1, 2000)}")


def new_workbook(instructions: list[str]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "안내"
    for line in instructions:
        ws.append([line])
    ws.column_dimensions["A"].width = 120
    return wb


def read_sheet(path: Path, title: str) -> list[dict[str, str]]:
    ws = load_workbook(path, read_only=True, data_only=True)[title]
    rows = ws.iter_rows(values_only=True)
    header = [str(h) for h in next(rows)]
    out = []
    for values in rows:
        record = {h: ("" if v is None else str(v).strip()) for h, v in zip(header, values)}
        if any(record.values()):
            out.append(record)
    return out
